"""Ordinal official result1 writer and disk-reopened independent readback."""
import argparse
import csv
from hashlib import sha256
import json
import math
from pathlib import Path
import openpyxl
from .io_utils import read_csv

def _hash(path): return sha256(Path(path).read_bytes()).hexdigest()

def _template_values(workbook):
    return [[cell.value for row in sheet.iter_rows() for cell in row] for sheet in workbook.worksheets]

def write_result1(template, schedule_path, output):
    schedule=read_csv(schedule_path)
    if len(schedule)!=144 or [r['slot_id'] for r in schedule]!=list(range(1,145)):
        raise ValueError('Validated 144-slot ordinal schedule required')
    wb=openpyxl.load_workbook(template)
    if len(wb.worksheets)!=2 or wb.worksheets[0].max_row!=145 or wb.worksheets[1].max_row!=7:
        raise ValueError('Unexpected official result1 layout')
    purchase,storage=wb.worksheets
    for t,row in enumerate(schedule,2): purchase.cell(t,2).value=row['grid_kwh']
    for block in range(6):
        rows=schedule[24*block:24*(block+1)]
        storage.cell(block+2,2).value=math.fsum(r['charge_kwh'] for r in rows)
        storage.cell(block+2,3).value=math.fsum(r['discharge_kwh'] for r in rows)
    storage['E2']=6000.0; storage['E3']=schedule[-1]['soc_end_kwh']
    Path(output).parent.mkdir(parents=True,exist_ok=True);wb.save(output);wb.close()

def readback(template, schedule_path, output, report_path, audit_path):
    schedule=read_csv(schedule_path)
    source=openpyxl.load_workbook(template,read_only=True,data_only=False)
    saved=openpyxl.load_workbook(output,read_only=True,data_only=False)
    sheet_names_unchanged=source.sheetnames==saved.sheetnames
    svals=_template_values(source);ovals=_template_values(saved)
    allowed={(0,r,2) for r in range(2,146)}|{(1,r,c) for r in range(2,8) for c in (2,3)}|{(1,2,5),(1,3,5)}
    unexpected=[]
    for si,(before,after) in enumerate(zip(svals,ovals)):
        width=source.worksheets[si].max_column
        for index,(a,b) in enumerate(zip(before,after)):
            row=index//width+1;col=index%width+1
            if (si,row,col) not in allowed and a!=b: unexpected.append(f'{si}:{row}:{col}')
    label_cells=[(0,r,1) for r in range(1,146)]+[(1,r,1) for r in range(1,8)]+[(1,r,4) for r in range(1,8)]
    labels_unchanged=all(source.worksheets[s].cell(r,c).value==saved.worksheets[s].cell(r,c).value for s,r,c in label_cells)
    rows=[];diffs=[]
    purchase=saved.worksheets[0];storage=saved.worksheets[1]
    for t,row in enumerate(schedule,1):
        actual=purchase.cell(t+1,2).value;diff=abs(float(actual)-row['grid_kwh']) if actual is not None else math.inf
        diffs.append(diff);rows.append(dict(record_type='slot_grid',ordinal=t,target_cell=f'B{t+1}',expected=row['grid_kwh'],actual=actual,abs_diff=diff,status='PASS' if diff<=1e-9 else 'FAIL'))
    aggregate_ok=True
    for block in range(6):
        segment=schedule[24*block:24*(block+1)]
        for col,key in ((2,'charge_kwh'),(3,'discharge_kwh')):
            expected=math.fsum(r[key] for r in segment);actual=storage.cell(block+2,col).value;diff=abs(float(actual)-expected) if actual is not None else math.inf
            aggregate_ok &= diff<=1e-9;rows.append(dict(record_type='4h_'+key,ordinal=block+1,target_cell=f'{chr(64+col)}{block+2}',expected=expected,actual=actual,abs_diff=diff,status='PASS' if diff<=1e-9 else 'FAIL'))
    soc_ok=storage['E2'].value==6000 and storage['E3'].value==6000
    for cell in ('E2','E3'): rows.append(dict(record_type='soc_endpoint',ordinal=1 if cell=='E2' else 145,target_cell=cell,expected=6000,actual=storage[cell].value,abs_diff=abs(float(storage[cell].value)-6000),status='PASS' if storage[cell].value==6000 else 'FAIL'))
    report={'status':'PASS' if sheet_names_unchanged and labels_unchanged and not unexpected and max(diffs)<=1e-9 and aggregate_ok and soc_ok else 'FAIL','mapping':'C_R1_RIGHT_ENDPOINT_ORDINAL_EXPORT','sheet_names_unchanged':sheet_names_unchanged,'sheet_names':saved.sheetnames,'purchase_values_present':sum(purchase.cell(r,2).value is not None for r in range(2,146)),'purchase_max_abs_diff_kwh':max(diffs),'aggregate_status':'PASS' if aggregate_ok else 'FAIL','soc_endpoint_status':'PASS' if soc_ok else 'FAIL','labels_unchanged':labels_unchanged,'unexpected_modified_cells':unexpected,'template_sha256':_hash(template),'output_sha256':_hash(output),'stage1_optimal_purchase_cost_cny':35126.948589289634,'stage2_export_schedule_cost_cny':math.fsum(r['cost_yuan'] for r in schedule),'stage2_cost_gap_cny':math.fsum(r['cost_yuan'] for r in schedule)-35126.948589289634,'stage2_cost_gap_bound_cny':0.0001,'result1_frozen':'PENDING_FYQ_CONTROLLER'}
    Path(report_path).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with Path(audit_path).open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    source.close();saved.close()
    if report['status']!='PASS': raise RuntimeError('RESULT1_READBACK_FAIL')
    return report

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--template',required=True);p.add_argument('--schedule',required=True);p.add_argument('--output',required=True);p.add_argument('--report',required=True);p.add_argument('--audit',required=True);a=p.parse_args()
    write_result1(a.template,a.schedule,a.output);print(json.dumps(readback(a.template,a.schedule,a.output,a.report,a.audit),ensure_ascii=False))
